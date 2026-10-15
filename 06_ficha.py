"""
Ficha de verificacion: XML -> Excel con forma de estado financiero
==================================================================

Genera, por fondo-anio, un Excel que se lee como el PDF, para comparar a
simple vista durante la verificacion manual:

  - una hoja por estado (ESF, ERI, ECP, EFE) con las cuentas en el orden
    del XML, nombres legibles y columnas Actual / Anterior;
  - el ECP con las series en columnas y el total al final (como el PDF);
  - formato de numero como el PDF: 1.234.567 / (1.234) / -;
  - por cada cifra: XML | PDF | Diferencia | Pagina. Si el cruce no
    encontro la cifra, se PROPONE el monto huerfano mas cercano del PDF;
  - en el periodo Anterior, ademas, el ORIGINAL: la cifra publicada en el
    XML del reporte del anio anterior, y su diferencia;
  - colores: rojo = no encontrado, naranja = signo distinto, amarillo =
    redondeo, morado = el comparativo difiere del original;
  - hoja Huerfanos: montos del PDF sin par en el XML (pagina y monto).

Uso:
  python 06_ficha.py <rut> <anio>   -> una ficha
  python 06_ficha.py --muestra      -> fichas de todos los casos de
                                       verificacion/muestra_verificacion.xlsx

Lee:      data/procesado/encabezados.csv, cuentas.csv, cruce_resumen.csv,
          cruce_detalle.csv, cruce_huerfanos.csv
Escribe:  verificacion/fichas/{rut}_{anio}.xlsx  (regenerables: no versionar)

Cada cuenta se asigna a su estado recorriendo el XML en orden: los codigos
exclusivos de un estado fijan la seccion y los compartidos (p.ej. diferencias
de cambio sobre efectivo, en ERI y EFE) heredan la seccion vigente.
"""

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
DIR_PROC = BASE_DIR / "data" / "procesado"
DIR_FICHAS = BASE_DIR / "verificacion" / "fichas"
MUESTRA = BASE_DIR / "verificacion" / "muestra_verificacion.xlsx"

FUENTE = Font(name="Arial", size=10)
FUENTE_B = Font(name="Arial", size=10, bold=True)
FMT_PDF = '#,##0;(#,##0);"-"'
COLORES = {
    "no_encontrado": PatternFill("solid", start_color="FFC7CE"),   # rojo
    "signo_distinto": PatternFill("solid", start_color="FFD8A8"),  # naranja
    "ok_redondeo": PatternFill("solid", start_color="FFF2A8"),     # amarillo
}
ESTADOS = ["ESF", "ERI", "ECP", "EFE"]
CONTEXTOS = {"PeriodoActual": "Actual", "PeriodoAnterior": "Anterior"}


# Codigos EXCLUSIVOS de cada estado (taxonomia CMF de fondos mutuos, en el
# orden del XML). Los que existen en dos estados (p.ej. diferencias de
# cambio sobre efectivo: ERI y EFE; distribucion de beneficios: ERI y ECP)
# o que no estan en la lista heredan la seccion en curso. Asi la
# clasificacion sigue el orden real del archivo, aunque una AGF ordene los
# estados distinto (p.ej. EFE antes que ECP).
ANCLAS = {
    "ESF": """EfectivoYEfectivoEquivalente ActivosFinancierosAValorRazonableConEfectoEnResultados
        ActivosFinancierosAValorRazonableConEfectoEnResultadosEntregadosEnGarantia
        ActivosFinancierosACostoAmortizado CuentasPorCobrarAIntermediarios OtrasCuentasPorCobrar
        OtrosActivos TotalActivo PasivosFinancierosAValorRazonableConEfectoEnResultados
        CuentasPorAPagarIntermediarios RescatesPorPagar RemuneracionesSociedadAdministradora
        OtrosDocumentosYCuentasPorPagar OtrosPasivos TotalPasivo ActivoNetoAtribuibleALosParticipes""",
    "ERI": """InteresesYReajustes IngresosPorDividendos
        DiferenciasDeCambioNetasSobreActivosFinancierosACostoAmortizado
        CambiosNetosEnValorRazonableDeActivosYPasivosFinancierosAValorRazonableConEfectoEnResultados
        ResultadoEnVentaDeInstrumentosFinancieros OtrosEri TotalIngresosPerdidasNetosDeLaOperacion
        ComisionDeAdministracion HonorariosPorCustodiaYAdministracion CostosDeTransaccion
        OtrosGastosDeOperacion TotalGastosDeOperacion UtilidadPerdidaDeLaOperacionAntesDeImpuesto
        ImpuestosALasGananciasPorInversionesEnElExterior UtilidadPerdidaDeLaOperacionDespuesDeImpuesto
        AumentoDisminucionDeActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionAntesDeDistribucionDeBeneficios
        AumentoDisminucionDeActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionDespuesDeDistribucionDeBeneficios""",
    "ECP": """ActivoNetoAtribuibleALosParticipesAl01DeEneroPorSerie ActivoNetoAtribuibleALosParticipesAl01DeEnero
        AportesDeCuotasPorSerie AportesDeCuotas RescateDeCuotasPorSerie RescateDeCuotas
        AumentoNetoOriginadoPorTransaccionesDeCuotasPorSerie AumentoNetoOriginadoPorTransaccionesDeCuotas
        AumentoDisminucionDelActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionAntesDeDistribucionDeBeneficiosPorSerie
        AumentoDisminucionDelActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionAntesDeDistribucionDeBeneficios
        DistribucionDeBeneficiosPorSerie DistribucionDeBeneficiosEnEfectivoPorSerie DistribucionDeBeneficiosEnEfectivo
        DistribucionDeBeneficiosEnCuotasPorSerie DistribucionDeBeneficiosEnCuotas
        AumentoDisminucionDelActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionDespuesDeDistribucionDeBeneficiosPorSerie
        AumentoDisminucionDelActivoNetoAtribuibleAParticipesOriginadasPorActividadesDeLaOperacionDespuesDeDistribucionDeBeneficios
        ActivoNetoAtribuibleAParticipesAl31DeDiciembreValorCuotaPorSerie ActivoNetoAtribuibleAParticipesAl31DeDiciembreValorCuota""",
    "EFE": """CompraDeActivosFinancieros VentaCobroDeActivosFinancieros InteresesDiferenciasDeCambioYReajustesRecibidos
        LiquidacionDeInstrumentosFinancierosDerivados DividendosRecibidos
        MontosPagadosASociedadAdministradoraEIntermediarios MontosRecibidosDeSociedadAdministradoraEIntermediarios
        OtrosIngresosDeOperacion OtrosGastosDeOperacionPagados FlujoNetoOriginadoPorActividadesDeLaOperacion
        ColocacionDeCuotasEnCirculacion RescatesDeCuotasEnCirculacion Otros
        FlujoNetoOriginadoPorActividadesDeFinanciamiento AumentoNetoDeEfectivoYEfectivoEquivalente
        SaldoInicialDeEfectivoYEfectivoEquivalente SaldoFinalDeEfectivoYEfectivoEquivalente""",
}
SECCION_DE = {c: est for est, txt in ANCLAS.items() for c in txt.split()}


def asignar_secciones(codigos):
    """Recorre los codigos en el orden del XML; un codigo exclusivo fija la
    seccion, los demas heredan la vigente."""
    actual, out = "ESF", []
    for c in codigos:
        actual = SECCION_DE.get(c, actual)
        out.append(actual)
    return out


def legible(codigo: str) -> str:
    """'VentaCobroDeActivosFinancieros' -> 'Venta cobro de activos financieros'."""
    palabras = re.findall(r"[A-Z][a-z0-9áéíóúñ]*|[a-z0-9]+|[A-Z]+(?![a-z])", codigo)
    texto = " ".join(palabras).lower()
    texto = re.sub(r"([a-z])(\d)", r"\1 \2", texto)          # "al01" -> "al 01"
    return texto[:1].upper() + texto[1:]


MAX_DIF_PROPUESTA = 0.10    # un huerfano se propone como "valor PDF" si difiere <= 10%
COLOR_ORIGINAL = PatternFill("solid", start_color="D9C3F0")   # morado: difiere del original t-1


class Datos:
    """Carga una sola vez lo necesario, para generar muchas fichas rapido.
    Incluye el XML del anio anterior de cada fondo, para la columna Original."""

    def __init__(self, claves):
        self.enc = pd.read_csv(DIR_PROC / "encabezados.csv", dtype={"rut": str, "rut_agf": str})
        self.enc = self.enc[self.enc["es_ultimo_envio"]]
        k = pd.DataFrame(claves, columns=["rut", "anio"])
        k_prev = k.assign(anio=k["anio"] - 1)
        archivos = pd.concat([self.enc.merge(k, on=["rut", "anio"]),
                              self.enc.merge(k_prev, on=["rut", "anio"])])["archivo"]
        cta = pd.read_csv(DIR_PROC / "cuentas.csv",
                          usecols=["archivo", "contexto", "codigo", "serie", "valor"],
                          dtype={"serie": str})
        self.cta = cta[cta["archivo"].isin(set(archivos))]
        self.res = pd.read_csv(DIR_PROC / "cruce_resumen.csv", dtype={"rut": str})
        self.det = pd.read_csv(DIR_PROC / "cruce_detalle.csv", dtype={"rut": str, "serie": str})
        self.hue = pd.read_csv(DIR_PROC / "cruce_huerfanos.csv", dtype={"rut": str})


class Comparador:
    """Arma, para cada cifra del XML, su contraparte en el PDF y en el original.

    El cruce sabe SI un monto esta en el PDF, no en QUE linea. Por eso:
      - encontrado igual      -> PDF = XML
      - signo distinto        -> PDF = -XML
      - redondeo              -> PDF = XML +-1 (no se guarda cual de los dos)
      - no encontrado         -> se PROPONE el huerfano mas cercano (<= 10%),
                                 con su pagina; si no hay, '?'
    La propuesta es una pista para mirar el PDF, no una lectura de la linea."""

    def __init__(self, det, hue, cta_original):
        self.estado, self.paginas = {}, {}
        for r in det.itertuples():
            k = (r.contexto, r.codigo, "" if pd.isna(r.serie) else r.serie)
            self.estado[k] = r.estado
            self.paginas[k] = "" if pd.isna(r.paginas) else str(r.paginas)
        self.huerfanos = list(hue.drop_duplicates("monto_pdf")[["monto_pdf", "pagina"]]
                              .itertuples(index=False, name=None))
        # Original: la columna "Actual" del reporte del anio anterior
        orig = cta_original[cta_original["contexto"] == "PeriodoActual"]
        self.original = {(c, "" if pd.isna(s) else s): v for c, s, v in
                         orig[["codigo", "serie", "valor"]].itertuples(index=False, name=None)}

    def pdf(self, ctx, codigo, serie, v):
        """Devuelve (valor_pdf, diferencia, pagina, estado)."""
        if v is None or pd.isna(v):
            return None, None, "", None
        k = (ctx, codigo, serie)
        est = self.estado.get(k)
        if est is None or est in ("ok", "ok_debil", "ok_union"):
            return v, 0, self.paginas.get(k, ""), None
        if est == "signo_distinto":
            return -v, -2 * v, self.paginas.get(k, ""), est
        if est == "ok_redondeo":
            return "±1", "±1", "", est
        if self.huerfanos and v:                      # no_encontrado: proponer
            h, pag = min(self.huerfanos, key=lambda x: abs(abs(x[0]) - abs(v)))
            if abs(abs(h) - abs(v)) <= MAX_DIF_PROPUESTA * abs(v):
                h = h if (h >= 0) == (v >= 0) else -h
                return h, h - v, f"{pag} (propuesto)", est
        return "?", None, "", est

    def orig(self, codigo, serie, v):
        """(valor_original, diferencia) o (None, None) si no hay reporte t-1."""
        if not self.original:
            return None, None
        o = self.original.get((codigo, serie))
        if o is None or v is None or pd.isna(v):
            return o, None
        return o, v - o


def _celda(ws, fila, col, v, fill=None, fmt=True):
    vacio = v is None or (not isinstance(v, str) and pd.isna(v))
    c = ws.cell(row=fila, column=col, value=None if vacio else v)
    c.font = FUENTE
    if isinstance(v, str):
        c.alignment = Alignment(horizontal="right")
    elif fmt:
        c.number_format = FMT_PDF
    if fill is not None:
        c.fill = fill
    return c


def _contextos(df):
    return [c for c in CONTEXTOS if c in set(df["contexto"])] + \
           sorted(set(df["contexto"].dropna()) - set(CONTEXTOS))


def hoja_estado(ws, cta_estado, comp: Comparador):
    """ESF / ERI / EFE: por periodo, XML | PDF | Dif | Pag (+ Original | Dif en Anterior)."""
    ctxs = _contextos(cta_estado)
    encab = ["Cuenta"]
    for ctx in ctxs:
        nombre = CONTEXTOS.get(ctx, ctx)
        encab += [f"{nombre} XML", f"{nombre} PDF", "Dif", "Pag"]
        if ctx == "PeriodoAnterior" and comp.original:
            encab += ["Original t-1", "Dif vs orig"]
    for j, t in enumerate(encab, 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = FUENTE_B
        c.alignment = Alignment(horizontal="center" if j > 1 else "left", wrap_text=True)

    fila = 1
    for r in cta_estado.drop_duplicates(["codigo", "serie"]).itertuples():
        fila += 1
        serie = "" if pd.isna(r.serie) else r.serie
        ws.cell(row=fila, column=1,
                value=legible(r.codigo) + (f"  [{serie}]" if serie else "")).font = FUENTE
        col = 2
        for ctx in ctxs:
            m = cta_estado[(cta_estado["codigo"] == r.codigo) & (cta_estado["contexto"] == ctx)
                           & (cta_estado["serie"].fillna("") == serie)]
            v = m["valor"].iloc[0] if len(m) else None
            p, dif, pag, est = comp.pdf(ctx, r.codigo, serie, v)
            fill = COLORES.get(est)
            _celda(ws, fila, col, v)
            _celda(ws, fila, col + 1, p, fill)
            _celda(ws, fila, col + 2, None if dif == 0 else dif, fill)
            _celda(ws, fila, col + 3, pag, fmt=False)
            col += 4
            if ctx == "PeriodoAnterior" and comp.original:
                o, d = comp.orig(r.codigo, serie, v)
                f2 = COLOR_ORIGINAL if d else None
                _celda(ws, fila, col, o, f2)
                _celda(ws, fila, col + 1, d if d else None, f2)
                col += 2
    ws.column_dimensions["A"].width = 60
    for j in range(2, len(encab) + 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 15
    ws.freeze_panes = "B2"


def hoja_ecp(ws, cta_ecp, comp: Comparador):
    """ECP: por periodo, bloques XML / PDF / Dif (y Original / Dif vs orig en
    Anterior), con las series en columnas y el Total al final."""
    cta_ecp = cta_ecp.copy()
    cta_ecp["base"] = cta_ecp["codigo"].str.replace("PorSerie$", "", regex=True)
    series = list(cta_ecp["serie"].dropna().drop_duplicates())
    cols = series + ["Total"]
    fila = 1
    for ctx in _contextos(cta_ecp):
        bloque = cta_ecp[cta_ecp["contexto"] == ctx]
        bases = list(bloque["base"].drop_duplicates())
        celdas = {}   # base -> [(xml, pdf, dif, est, orig, dif_orig) por columna]
        for base in bases:
            g = bloque[bloque["base"] == base]
            fila_c = []
            for s in series + [None]:
                m = g[g["serie"] == s] if s else g[g["serie"].isna()]
                if not len(m):
                    fila_c.append((None,) * 6)
                    continue
                cod, v = m.iloc[0]["codigo"], m.iloc[0]["valor"]
                p, dif, _, est = comp.pdf(ctx, cod, s or "", v)
                o, do = comp.orig(cod, s or "", v)
                fila_c.append((v, p, dif, est, o, do))
            celdas[base] = fila_c

        vistas = [("XML", 0, None), ("PDF", 1, "cruce"), ("Diferencia PDF - XML", 2, "cruce")]
        if ctx == "PeriodoAnterior" and comp.original:
            vistas += [("Original t-1 (XML del reporte anterior)", 4, None),
                       ("Diferencia vs original", 5, "orig")]
        nombre = CONTEXTOS.get(ctx, ctx)
        for titulo, idx, color in vistas:
            ws.cell(row=fila, column=1, value=f"{nombre} - {titulo}").font = FUENTE_B
            fila += 1
            ws.cell(row=fila, column=1, value="Cuenta").font = FUENTE_B
            for j, t in enumerate(cols, 2):
                c = ws.cell(row=fila, column=j, value=t)
                c.font = FUENTE_B
                c.alignment = Alignment(horizontal="center")
            for base in bases:
                fila += 1
                ws.cell(row=fila, column=1, value=legible(base)).font = FUENTE
                for j, t in enumerate(celdas[base], 2):
                    v = t[idx]
                    if idx in (2, 5) and v == 0:
                        v = None
                    if color == "orig":
                        fill = COLOR_ORIGINAL if t[5] else None
                    elif color == "cruce":
                        fill = COLORES.get(t[3])
                    else:
                        fill = None
                    _celda(ws, fila, j, v, fill)
            fila += 2
    ws.column_dimensions["A"].width = 60
    for j in range(2, len(cols) + 2):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 14


def generar(datos: Datos, rut: str, anio: int):
    enc_t = datos.enc[(datos.enc["rut"] == rut) & (datos.enc["anio"] == anio)]
    if enc_t.empty:
        print(f"  ! {rut} {anio}: no esta en encabezados.csv")
        return None
    enc = enc_t.iloc[0]
    cta = datos.cta[datos.cta["archivo"] == enc["archivo"]].copy()
    cta = cta[cta["valor"].notna()]
    cta["estado_fin"] = asignar_secciones(cta["codigo"].tolist())   # orden del XML
    prev = datos.enc[(datos.enc["rut"] == rut) & (datos.enc["anio"] == anio - 1)]
    cta_orig = datos.cta[datos.cta["archivo"].isin(prev["archivo"])]
    det = datos.det[(datos.det["rut"] == rut) & (datos.det["anio"] == anio)]
    hue = datos.hue[(datos.hue["rut"] == rut) & (datos.hue["anio"] == anio)]
    res = datos.res[(datos.res["rut"] == rut) & (datos.res["anio"] == anio)]
    comp = Comparador(det, hue, cta_orig)

    wb = Workbook()
    info = wb.active
    info.title = "Fondo"
    pdf = res["pdf"].iloc[0] if len(res) and pd.notna(res["pdf"].iloc[0]) else ""
    pares = [
        ("Fondo", enc["nombre_fondo"]), ("RUT", rut), ("Anio", anio),
        ("AGF", enc["nombre_agf"]), ("Auditor", enc.get("auditor_firma")),
        ("Moneda", enc["moneda"]), ("XML", enc["archivo"]),
        ("XML original (t-1)", prev["archivo"].iloc[0] if len(prev) else "(no disponible)"),
        ("PDF", str(Path("data/archivos") / rut / str(anio) / pdf) if pdf else ""),
        ("Paginas de estados en el PDF", res["paginas_estados"].iloc[0] if len(res) else ""),
        ("% encontrado (tolerante)", res["pct_tolerante"].iloc[0] if len(res) else ""),
        ("", ""),
        ("Como leer", ""),
        ("  PDF", "valor del PDF segun el cruce. Si no se encontro, se PROPONE el monto huerfano "
                  "mas cercano (<=10%), con su pagina: es una pista, confirmar en el PDF"),
        ("  Original t-1", "columna Actual del XML del reporte anterior (solo periodo Anterior)"),
        ("  rojo", "cifra del XML no encontrada en el PDF (PDF = propuesta o '?')"),
        ("  naranja", "monto encontrado con signo distinto"),
        ("  amarillo", "encontrado con diferencia de redondeo (+-1)"),
        ("  morado", "el comparativo difiere del original publicado el anio anterior"),
    ]
    for i, (k, v) in enumerate(pares, 1):
        info.cell(row=i, column=1, value=k).font = FUENTE_B
        c = info.cell(row=i, column=2, value=None if pd.isna(v) else v)
        c.font = FUENTE
        c.alignment = Alignment(wrap_text=True)
    etiquetas = [p[0] for p in pares]
    for k, fill in (("  rojo", COLORES["no_encontrado"]), ("  naranja", COLORES["signo_distinto"]),
                    ("  amarillo", COLORES["ok_redondeo"]), ("  morado", COLOR_ORIGINAL)):
        info.cell(row=etiquetas.index(k) + 1, column=1).fill = fill
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 90

    for est in ESTADOS:
        ws = wb.create_sheet(est)
        parte = cta[cta["estado_fin"] == est]
        if parte.empty:
            ws["A1"] = "(sin cuentas en el XML para este estado)"
            continue
        (hoja_ecp if est == "ECP" else hoja_estado)(ws, parte, comp)

    hu = wb.create_sheet("Huerfanos")
    hu["A1"], hu["B1"] = "Pagina PDF", "Monto en PDF sin par en el XML"
    hu["A1"].font = hu["B1"].font = FUENTE_B
    for i, r in enumerate(hue.drop_duplicates(["pagina", "monto_pdf"])
                          .sort_values(["pagina", "monto_pdf"]).itertuples(), 2):
        hu.cell(row=i, column=1, value=r.pagina).font = FUENTE
        c = hu.cell(row=i, column=2, value=r.monto_pdf)
        c.font, c.number_format = FUENTE, FMT_PDF
    hu.column_dimensions["A"].width = 12
    hu.column_dimensions["B"].width = 32

    DIR_FICHAS.mkdir(parents=True, exist_ok=True)
    ruta = DIR_FICHAS / f"{rut}_{anio}.xlsx"
    wb.save(ruta)
    return ruta


def main():
    if len(sys.argv) == 2 and sys.argv[1] == "--muestra":
        m = pd.read_excel(MUESTRA, sheet_name="Muestra", dtype={"rut": str})
        claves = list(zip(m["rut"], m["anio"].astype(int)))
    elif len(sys.argv) == 3:
        claves = [(sys.argv[1], int(sys.argv[2]))]
    else:
        print(__doc__)
        return 1

    datos = Datos(claves)
    for rut, anio in claves:
        ruta = generar(datos, rut, anio)
        if ruta:
            print(f"  {rut} {anio} -> {ruta.relative_to(BASE_DIR)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())