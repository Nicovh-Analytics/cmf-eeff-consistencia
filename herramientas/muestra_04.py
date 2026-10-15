"""
Muestra de verificacion manual
==============================

Arma una planilla Excel con una muestra estratificada de candidatos para
revisar PDF en mano y anotar si cada uno es una discrepancia real.

Diseno de la muestra (por cada tipo de candidato):
  - N_TOP "mas_material": los de mayor diferencia relativa al activo neto.
    Sirven para conocer el IMPACTO de los casos grandes.
  - N_AZAR "aleatorio": al azar entre el resto, con semilla fija.
    Son los que permiten estimar la PRECISION del metodo (que fraccion de
    los candidatos de ese tipo es real). Los "mas_material" NO entran en
    esa estimacion: estan seleccionados a proposito, no al azar.

NUNCA sobrescribe: si la planilla ya existe, no hace nada. Tus anotaciones
son trabajo que no se puede regenerar.

Lee:      data/procesado/candidatos.csv, cruce_resumen.csv,
          cruce_detalle.csv, cruce_huerfanos.csv
Escribe:  verificacion/muestra_verificacion.xlsx   (fuera del repo: .gitignore)
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = Path(__file__).resolve().parent.parent   # raiz del repo (este script vive en herramientas/)
DIR_PROC = BASE_DIR / "data" / "procesado"
DIR_ARCHIVOS = BASE_DIR / "data" / "archivos"
SALIDA = BASE_DIR / "verificacion" / "muestra_verificacion.xlsx"

N_TOP = 5
N_AZAR = 5
SEMILLA = 42
MAX_ITEMS = 6            # cifras faltantes / huerfanos a mostrar por caso
RESULTADOS = ["confirmado", "falso_positivo", "dudoso"]

FUENTE = Font(name="Arial", size=10)
FUENTE_TITULO = Font(name="Arial", size=10, bold=True)
RELLENO_EDITAR = PatternFill("solid", start_color="FFFF00")   # celdas a llenar


def armar_muestra(cand: pd.DataFrame) -> pd.DataFrame:
    partes = []
    for tipo, grupo in cand.groupby("tipo"):
        top = grupo.sort_values("materialidad_pct", ascending=False, na_position="last").head(N_TOP)
        resto = grupo.drop(top.index)
        azar = resto.sample(n=min(N_AZAR, len(resto)), random_state=SEMILLA)
        partes += [top.assign(seleccion="mas_material"), azar.assign(seleccion="aleatorio")]
    return pd.concat(partes, ignore_index=True)


def contexto_de_cada_caso(muestra, det, hue):
    """Texto con las cifras faltantes y los huerfanos de cada caso, para
    no tener que buscarlos en los CSV mientras se revisa el PDF."""
    claves = muestra[["rut", "anio"]]
    falt = det[det["estado"] == "no_encontrado"].merge(claves, on=["rut", "anio"])
    hue = hue.merge(claves, on=["rut", "anio"])

    def fmt_falt(g):
        g = g.reindex(g["valor_xml"].abs().sort_values(ascending=False).index).head(MAX_ITEMS)
        return "\n".join(f"{r.codigo} [{str(r.contexto).replace('Periodo', '')}"
                         f"{'' if pd.isna(r.serie) else ' / ' + str(r.serie)}] {r.valor_xml:,.0f}"
                         .replace(",", ".") for r in g.itertuples())

    def fmt_hue(g):
        g = g.drop_duplicates("monto_pdf")
        g = g.reindex(g["monto_pdf"].abs().sort_values(ascending=False).index).head(MAX_ITEMS)
        return "\n".join(f"pag {r.pagina}: {r.monto_pdf:,.0f}".replace(",", ".")
                         for r in g.itertuples())

    tf = falt.groupby(["rut", "anio"]).apply(fmt_falt, include_groups=False).rename("faltantes_xml")
    th = hue.groupby(["rut", "anio"]).apply(fmt_hue, include_groups=False).rename("huerfanos_pdf_detalle")
    return muestra.merge(tf, on=["rut", "anio"], how="left").merge(th, on=["rut", "anio"], how="left")


def escribir_excel(m: pd.DataFrame):
    columnas = [
        ("id", 6), ("seleccion", 13), ("tipo", 22), ("rut", 8), ("anio", 6),
        ("nombre_fondo", 30), ("nombre_agf", 28), ("auditor_firma", 12),
        ("delta_principal", 14), ("materialidad_pct", 13), ("estados", 10),
        ("periodos", 14), ("paginas_estados", 11), ("ruta_pdf", 40),
        ("faltantes_xml", 55), ("huerfanos_pdf_detalle", 26),
        ("resultado", 15), ("nota", 45),
    ]
    wb = Workbook()

    # --- Hoja Muestra ---
    ws = wb.active
    ws.title = "Muestra"
    for j, (nombre, ancho) in enumerate(columnas, 1):
        c = ws.cell(row=1, column=j, value=nombre)
        c.font = FUENTE_TITULO
        ws.column_dimensions[c.column_letter].width = ancho
    for i, fila in enumerate(m.itertuples(index=False), 2):
        for j, (nombre, _) in enumerate(columnas, 1):
            v = getattr(fila, nombre, None)
            c = ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
            c.font = FUENTE
            c.alignment = Alignment(vertical="top", wrap_text=nombre in
                                    ("faltantes_xml", "huerfanos_pdf_detalle", "nota", "nombre_fondo"))
            if nombre == "delta_principal":
                c.number_format = "#,##0"
            if nombre in ("resultado", "nota"):
                c.fill = RELLENO_EDITAR
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions

    col_res = [n for n, _ in columnas].index("resultado") + 1
    letra_res = ws.cell(row=1, column=col_res).column_letter
    dv = DataValidation(type="list", formula1=f'"{",".join(RESULTADOS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letra_res}2:{letra_res}{len(m) + 1}")

    # --- Hoja Resumen (formulas: se actualiza sola al anotar) ---
    rs = wb.create_sheet("Resumen")
    rs["A1"] = "Precision estimada por tipo (solo filas 'aleatorio')"
    rs["A1"].font = FUENTE_TITULO
    encab = ["tipo", "en muestra", "revisados", "confirmado", "falso_positivo", "dudoso", "precision"]
    for j, t in enumerate(encab, 1):
        rs.cell(row=3, column=j, value=t).font = FUENTE_TITULO
    letra_sel = ws.cell(row=1, column=[n for n, _ in columnas].index("seleccion") + 1).column_letter
    letra_tipo = ws.cell(row=1, column=[n for n, _ in columnas].index("tipo") + 1).column_letter
    rango = lambda L: f"Muestra!${L}$2:${L}${len(m) + 1}"
    for i, tipo in enumerate(sorted(m["tipo"].unique()), 4):
        rs.cell(row=i, column=1, value=tipo)
        base = f'{rango(letra_tipo)},$A{i},{rango(letra_sel)},"aleatorio"'
        rs.cell(row=i, column=2, value=f"=COUNTIFS({base})")
        for j, res in enumerate(RESULTADOS, 4):
            rs.cell(row=i, column=j, value=f'=COUNTIFS({base},{rango(letra_res)},"{res}")')
        rs.cell(row=i, column=3, value=f"=SUM(D{i}:F{i})")
        c = rs.cell(row=i, column=7, value=f'=IF(D{i}+E{i}=0,"",D{i}/(D{i}+E{i}))')
        c.number_format = "0%"
    nota = len(m["tipo"].unique()) + 5
    rs.cell(row=nota, column=1, value=(
        "Precision = confirmados / (confirmados + falsos positivos). Los 'dudoso' no entran. "
        f"Con {N_AZAR} casos por tipo es una estimacion gruesa: sirve para ordenar, no para concluir."))
    for fila in rs.iter_rows():
        for c in fila:
            if c.font != FUENTE_TITULO:
                c.font = FUENTE
    rs.column_dimensions["A"].width = 26

    # --- Hoja Leyenda ---
    ly = wb.create_sheet("Leyenda")
    lineas = [
        ("Como usar esta planilla", True),
        ("Solo se editan las columnas amarillas de la hoja Muestra: 'resultado' (lista desplegable) y 'nota'.", False),
        ("resultado = confirmado: el PDF y el XML muestran cifras distintas para la misma cuenta.", False),
        ("resultado = falso_positivo: las cifras coinciden; el candidato fue un problema de lectura del script.", False),
        ("resultado = dudoso: no se puede determinar (p.ej. el PDF no muestra esa linea).", False),
        ("En 'nota': cuenta, valor en XML, valor en PDF, pagina, y cualquier observacion.", False),
        ("", False),
        ("Ejemplo de fila anotada (caso del analisis exploratorio):", True),
        ("resultado: confirmado", False),
        ("nota: EFE actual, compras XML -87.353.981 / PDF (87.366.218), pag 10. "
         "Misma diferencia (12.237) en flujo neto de operacion y aumento neto de efectivo.", False),
        ("", False),
        ("No modificar las demas columnas: vienen del pipeline (04_candidatos.py).", False),
        ("La muestra 'aleatorio' usa semilla fija; volver a correr herramientas/muestra_04.py no sobrescribe este archivo.", False),
    ]
    for i, (texto, titulo) in enumerate(lineas, 1):
        c = ly.cell(row=i, column=1, value=texto)
        c.font = FUENTE_TITULO if titulo else FUENTE
    ly.column_dimensions["A"].width = 120

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)


def main():
    if SALIDA.exists():
        print(f"Ya existe {SALIDA}. No se sobrescribe (contiene tus anotaciones).")
        print("Si quieres una muestra nueva, renombra o mueve el archivo actual primero.")
        return 0

    cand = pd.read_csv(DIR_PROC / "candidatos.csv", dtype={"rut": str, "rut_agf": str})
    res = pd.read_csv(DIR_PROC / "cruce_resumen.csv", dtype={"rut": str})
    det = pd.read_csv(DIR_PROC / "cruce_detalle.csv", dtype={"rut": str, "serie": str})
    hue = pd.read_csv(DIR_PROC / "cruce_huerfanos.csv", dtype={"rut": str})

    m = armar_muestra(cand)
    m = m.merge(res[["rut", "anio", "paginas_estados"]], on=["rut", "anio"], how="left")
    m["ruta_pdf"] = [str(Path("data/archivos") / r / str(a) / p)
                     for r, a, p in zip(m["rut"], m["anio"], m["pdf"])]
    m = contexto_de_cada_caso(m, det, hue)
    m.insert(0, "id", range(1, len(m) + 1))
    m["resultado"] = None
    m["nota"] = None

    escribir_excel(m)
    print(f"Muestra: {len(m)} casos")
    print(m.groupby(["tipo", "seleccion"]).size().to_string())
    print(f"\nsalida: {SALIDA}")
    return 0


if __name__ == "__main__":
    sys.exit(main())